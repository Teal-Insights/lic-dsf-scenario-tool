"""Portable briefing files from one explicitly shareable comparison snapshot."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from zipfile import ZipFile, ZIP_DEFLATED

from .charts import METRICS, comparison_summary, render_comparison, render_indicator
from .ida21 import INPUTS
from .reader_guide import reader_sections, technical_glossary, financing_comparison_note, comparison_orientation, guidance_references


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, indent=2).encode('utf-8')


def _safe_text(value):
    # CSV consumers may treat user-controlled labels/notes as formulas.
    # Numeric cells remain numeric; JSON retains the exact original strings.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def _csv(rows):
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerows([_safe_text(v) for v in row] for row in rows)
    return stream.getvalue().encode('utf-8-sig')


def tables(comparison):
    summary = comparison_summary(comparison)
    evidence = 'Calculated by this tool; fresh Excel verification pending. Selected workbook years only.'
    disclosure = comparison.get('workbook_provenance', {}).get('disclosure', 'Analyst-supplied workbook.')
    sha = comparison['workbook_sha256']
    common_headers = ['workbook_sha256', 'evidence', 'input_disclosure']
    common = [sha, evidence, disclosure]
    results = [['indicator', 'metric', 'year', 'units', 'scenario', 'comparator',
                'reference_baseline', 'as_supplied_customized', 'scenario_value',
                'difference_pp', 'threshold', 'baseline_source_cell', 'customized_source_cell'] + common_headers]
    titles = {m[0]: m[1] for m in METRICS}
    for row in summary['rows']:
        for label in summary['labels']:
            results.append([titles[row['metric']], row['metric'], row['year'], row['units'],
                label, summary['comparator'], row['reference_baseline'], row['as_supplied_customized'],
                row['scenarios'][label], row['differences'][label], row['threshold'],
                row['source_cells'].get('reference_baseline'), row['source_cells'].get('customized')] + common)
    inputs = [['scenario', 'input_row', 'projection_year', 'adjustment', 'input', 'units'] + common_headers]
    terms = [['scenario', 'treatment', 'interest_rate_percent', 'grace_years', 'maturity_years'] + common_headers]
    notes = [['scenario', 'input', 'changed', 'shared_explanation'] + common_headers]
    for run, assumption in zip(comparison['runs'], summary['assumptions']):
        years = run['result']['input_years']
        for row, values in assumption['delta_paths'].items():
            if len(values) != len(years):
                raise ValueError('Input years and adjustments do not align.')
            driver = next(item for item in INPUTS if str(item[0]) == row)
            for year, value in zip(years, values):
                inputs.append([run['share_label'], row, year, value, driver[3], driver[4]] + common)
        financing = assumption['financing_terms']
        if isinstance(financing, dict):
            terms.append([run['share_label'], 'Scenario override', financing.get('rate', 0) * 100,
                          financing.get('grace_years'), financing.get('maturity_years')] + common)
        else:
            terms.append([run['share_label'], 'As supplied in workbook', None, None, None] + common)
        for entry in assumption['shared_explanations']:
            notes.append([run['share_label'], entry['label'], entry['changed'], entry['text']] + common)
    evidence_rows = [['scenario', 'revision', 'result_sha256', 'scenario_sha256',
                      'contract_version', 'engine_identity_json', 'evidence_json', 'warnings_json'] + common_headers]
    for run in comparison['runs']:
        r = run['result']
        evidence_rows.append([run['share_label'], run['revision'], run['result_hash'], r['scenario_hash'],
            r['contract_version'], json.dumps(r['engine_identity'], sort_keys=True),
            json.dumps(r['evidence'], sort_keys=True), json.dumps(r['warnings'], ensure_ascii=False)] + common)
    return {'Results': results, 'Macro assumptions': inputs, 'Financing': terms,
            'Shared explanations': notes, 'Evidence': evidence_rows}


def workbook_bytes(table_map, comparison=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    orientation = comparison_orientation(comparison_summary(comparison)) if comparison is not None else [('Source and comparison', 'See the scenario, comparator, workbook fingerprint and input disclosure columns in Results for this table export.')]
    guide = [['Read me first', 'How to use these tables']] + [list(item) for item in orientation + reader_sections()] + [[title, text] for title, text, url in guidance_references()] + [['Technical field', 'Meaning']] + [list(item) for item in technical_glossary()]
    table_map = {'Read me first': guide, **table_map}
    wb.remove(wb.active)
    wb.properties.creator = 'LIC-DSF Scenario Analysis Tool'
    wb.properties.title = 'Scenario comparison: data, assumptions and evidence'
    for name, rows in table_map.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    # Explicit strings also preserve '#N/A' as data, not an Excel error.
                    cell.data_type = 's'
                cell.font = Font(name='Calibri', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = '0.###############'
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='143E5A')
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for index, title in enumerate(rows[0], 1):
            ws.column_dimensions[get_column_letter(index)].width = 48 if any(x in title for x in ('explanation', 'evidence', 'identity', 'disclosure')) else 26
    guide_sheet = wb['Read me first']
    guide_sheet.column_dimensions['A'].width = 36
    guide_sheet.column_dimensions['B'].width = 105
    guide_sheet.auto_filter.ref = None
    links = {title: url for title, text, url in guidance_references()}
    if comparison is not None and comparison_summary(comparison)['provenance']['kind'] == 'official_illustrative_template':
        links['Starting source'] = 'https://thedocs.worldbank.org/en/doc/f0ade6bcf85b6f98dbeb2c39a2b7770c-0360012025/new-lic-dsf-template'
    for row in guide_sheet.iter_rows(min_row=2):
        if row[0].value in links:
            row[1].hyperlink = links[row[0].value]
            row[1].font = Font(name='Calibri', size=11, color='143E5A', underline='single')
    for row in guide_sheet.iter_rows(min_row=2):
        guide_sheet.row_dimensions[row[0].row].height = max(32, 16 * (len(str(row[1].value)) // 95 + 2))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pack(comparison):
    """No workspace access: callers must pass the existing shareable allowlist."""
    table_map = tables(comparison)
    files = {'data-and-assumptions.json': _json(comparison),
             'tables.xlsx': workbook_bytes(table_map, comparison)}
    for name, rows in table_map.items():
        files['tables/' + name.lower().replace(' ', '-') + '.csv'] = _csv(rows)
    for view, label in [('standard', 'standard'), ('briefing', 'policy')]:
        files[f'reports/{label}-report.pdf'] = render_comparison(comparison, view, 'pdf')
        for metric, *_ in METRICS:
            for fmt in ('png', 'svg'):
                files[f'charts/{label}/{metric}.{fmt}'] = render_indicator(comparison, metric, view, fmt)
    reading = ['START HERE: LIC-DSF scenario briefing', '',
        'Open reports/policy-report.pdf for the reading guide and comparison.',
        'Open tables.xlsx, starting with Read me first, to inspect or reuse the numbers.', '']
    for heading, body in comparison_orientation(comparison_summary(comparison)) + reader_sections():
        reading.extend([heading.upper(), body, ''])
    if comparison_summary(comparison)['provenance']['kind'] == 'official_illustrative_template':
        reading.extend(['Official starting-template source: https://thedocs.worldbank.org/en/doc/f0ade6bcf85b6f98dbeb2c39a2b7770c-0360012025/new-lic-dsf-template', ''])
    reading.append('SPECIFIC GUIDANCE REFERENCES (printed page numbers)')
    for title, text, url in guidance_references():
        reading.extend([title + ': ' + text, url, ''])
    reading.extend(['FILES IN THIS PACKET',
        'reports/: Standard and Policy PDFs, each with a reading guide and shared explanations.',
        'charts/: separate PNG images and scalable SVG figures for all six indicators in both styles.',
        'tables.xlsx: reader guide plus results, assumptions, financing, explanations and evidence.',
        'tables/*.csv: the same five data tables for reuse in other software.',
        'data-and-assumptions.json: full exported records and original shared text for technical reuse.',
        'manifest.json: file sizes and fingerprints for checking later changes.', '',
        'Values in CSV/JSON are not presentation-rounded; Excel uses its native numeric precision.',
        'Blank result cells are missing observations, never zero. Error strings remain strings.',
        'CSV text that could be treated as a formula has a protective apostrophe; Excel uses explicit text cells.',
        'The original text remains in JSON. This packet is not an importable scenario file or workspace backup.', '',
        'TECHNICAL FIELD GLOSSARY'])
    for field, meaning in technical_glossary():
        reading.append(field + ': ' + meaning)
    files['READ ME.txt'] = ('\n'.join(reading) + '\n').encode('utf-8')
    manifest = {'format': 'lic-dsf-briefing-packet-v1', 'workbook_sha256': comparison['workbook_sha256'],
        'comparator_id': comparison['comparator_id'], 'provenance': comparison.get('workbook_provenance'),
        'files': {name: {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()}
                  for name, data in sorted(files.items())}}
    files['manifest.json'] = _json(manifest)
    out = io.BytesIO()
    with ZipFile(out, 'w', compression=ZIP_DEFLATED) as archive:
        for name, data in sorted(files.items()):
            archive.writestr(name, data)
    return out.getvalue()
