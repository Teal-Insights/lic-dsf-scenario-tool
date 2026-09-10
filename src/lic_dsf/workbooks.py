"""Read-only package intake. No workbook conversion or external-link refresh."""
from __future__ import annotations

import hashlib
import math
import re
from pathlib import Path, PurePosixPath
import stat
import zipfile

from .contracts import CONTRACT_VERSION
from types import SimpleNamespace
from .ida21 import (PUBLIC, EXTERNAL, EXTRA_SHEETS, INPUTS, TERMS, METRICS,
                    THRESHOLDS, probes, cell)

MAX_ARCHIVE = 32 * 1024 * 1024
MAX_EXPANDED = 256 * 1024 * 1024
MAX_PART = 64 * 1024 * 1024
MAX_MEMBERS = 4000


class IntakeError(ValueError):
    """A stable error code, without a local path or spreadsheet contents."""


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _number(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value))


def _contains(value, expected):
    return isinstance(value, str) and expected in value.strip().lower()


def _additive_input(formula, address):
    """Accept scalar reference arithmetic only; functions may distinguish blanks."""
    if not isinstance(formula, str) or not formula.startswith('='):
        return False
    compact = re.sub(r'\s+', '', formula).replace('$', '').upper()
    reference = r"'[^']+'![A-Z]+[1-9][0-9]*"
    expression = r'=\+?' + reference + r'(?:[+-]\+?' + reference + r')*\+' + re.escape(address.upper())
    return re.fullmatch(expression, compact) is not None


def check_archive(path):
    path = Path(path)
    if not path.is_file():
        raise IntakeError("workbook_missing")
    if path.suffix.lower() not in {'.xlsm', '.xlsx'}:
        raise IntakeError("unsupported_file_extension")
    if path.stat().st_size > MAX_ARCHIVE:
        raise IntakeError("archive_size_limit")
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_MEMBERS or sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise IntakeError("expanded_size_limit")
            names = set()
            for info in infos:
                name = info.filename
                parts = PurePosixPath(name).parts
                if (name.startswith('/') or '\\' in name or ':' in name
                        or '..' in parts or '.' in parts or '\x00' in name
                        or name.casefold() in names):
                    raise IntakeError("unsafe_archive_path")
                names.add(name.casefold())
                if stat.S_ISLNK(info.external_attr >> 16) or info.flag_bits & 1:
                    raise IntakeError("unsupported_archive_entry")
                if info.file_size > MAX_PART or info.file_size > max(1, info.compress_size) * 1000:
                    raise IntakeError("expanded_part_limit")
                if name.endswith('.xml') or name.endswith('.rels'):
                    content = archive.read(info)
                    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                        raise IntakeError("xml_declaration_not_supported")
            if not {'[content_types].xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels'} <= names:
                raise IntakeError("invalid_workbook_package")
            return {'macro_part_present': 'xl/vbaproject.bin' in names,
                    'external_links_present': any(n.startswith('xl/externallinks/') for n in names)}
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise IntakeError("invalid_workbook_package") from exc


class _Grid:
    def __init__(self, sheet):
        self.cells = {(c.row, c.column): c for row in sheet.iter_rows(min_row=1, max_row=100, max_col=25)
                      for c in row if hasattr(c, 'row')}
    def cell(self, row, column):
        return self.cells.get((row, column), SimpleNamespace(value=None, data_type='n', coordinate=cell(column, row)))
    def __getitem__(self, address):
        from openpyxl.utils.cell import coordinate_to_tuple
        return self.cell(*coordinate_to_tuple(address))


class _Book:
    def __init__(self, workbook, names):
        self.original = workbook
        self.sheetnames = workbook.sheetnames
        self.grids = {name: _Grid(workbook[name]) for name in names if name in self.sheetnames}
    def __getitem__(self, name):
        return self.grids[name]
    def close(self):
        self.original.close()


def inspect_workbook(path):
    """Return serializable structural facts; no engine or Excel is invoked."""
    import openpyxl
    path = Path(path)
    package = check_archive(path)
    original_hash = file_hash(path)
    errors, warnings = [], []
    if package['macro_part_present']:
        warnings.append('macros_present_not_executed')
    if package['external_links_present']:
        warnings.append('external_links_present_not_refreshed')
    try:
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise IntakeError('workbook_parse_failed') from exc
    needed = {PUBLIC, EXTERNAL, 'Input 1 - Basics', 'Macro-Debt_Data', 'CI Summary'} | {p['sheet'] for p in probes()}
    formulas = _Book(formulas, needed)
    cached = _Book(cached, needed)
    result = {'contract_version': CONTRACT_VERSION, 'workbook_sha256': original_hash,
              'compatibility': {'status': 'unsupported', 'profile': 'ida21-fixed-english',
                                'checks': [], 'errors': errors},
              'first_projection_year': None, 'input_years': [], 'input_specs': [],
              'customization': {'delta_paths': {}, 'formula_cells': [], 'blank_cells': [],
                                'terms': {}, 'has_nonzero_deltas': False,
                                'blank_normalization': 'direct_additive_empty_cell_to_zero'},
              'thresholds': {}, 'warnings': warnings, 'package': package,
              'evidence': {'structural': 'unchecked', 'saved_cache': 'unchecked',
                           'exact_excel': 'pending', 'formula_profile': 'not_verified'}}
    try:
        required = set(EXTRA_SHEETS) | {PUBLIC, EXTERNAL, 'Input 1 - Basics', 'Macro-Debt_Data'}
        required |= {m[3] for m in METRICS} | {p['sheet'] for p in probes()}
        if not required <= set(formulas.sheetnames):
            errors.append('required_template_sheet_missing')
            return result
        basics = cached['Input 1 - Basics']
        first = basics['C18'].value
        if not _contains(basics['B18'].value, 'first year of projection') or not _number(first) or int(first) != first or not 1990 <= first <= 2100:
            errors.append('first_projection_year_invalid')
        else:
            result['first_projection_year'] = int(first)
        years = [cached[PUBLIC].cell(7, c).value for c in range(5, 26)]
        if result['first_projection_year'] is None or years != [result['first_projection_year'] + i for i in range(21)]:
            errors.append('input_year_geometry_invalid')
        else:
            result['input_years'] = [int(y) for y in years]
        for row, level, fragment, title, units in INPUTS:
            if not _contains(cached[PUBLIC].cell(level, 2).value, fragment) or str(cached[PUBLIC].cell(row, 2).value).strip() != 'Delta':
                errors.append(f'input_anchor_invalid:{row}')
            if row == 17:
                currency, scale = cached['Macro-Debt_Data']['C50'].value, cached['Macro-Debt_Data']['D50'].value
                if currency != 'National Currency' or scale not in {'Million', 'Billion'}:
                    errors.append('asset_units_unverified')
                units = f'local currency ({scale.lower()})' if isinstance(scale, str) else 'local currency (unverified scale)'
            result['input_specs'].append({'row': row, 'name': title, 'units': units})
            values = []
            for col in range(5, 26):
                f, v = formulas[PUBLIC].cell(row, col), cached[PUBLIC].cell(row, col).value
                additive = _additive_input(formulas[PUBLIC].cell(level, col).value, f.coordinate)
                if not additive:
                    errors.append(f'input_formula_geometry_invalid:{row}')
                if f.data_type == 'f':
                    result['customization']['formula_cells'].append(f.coordinate)
                if v is None:
                    result['customization']['blank_cells'].append(f.coordinate)
                    if f.data_type != 'f' and f.value is None and additive:
                        values.append(0.0)
                    else:
                        values.append(None)
                        errors.append(f'input_cache_or_blank_semantics_unverified:{row}')
                elif not _number(v):
                    errors.append(f'input_value_invalid:{row}')
                    values.append(None)
                else:
                    values.append(float(v))
            result['customization']['delta_paths'][str(row)] = values
        # Formula-valued delta cells need an explicit import design, not cached-value substitution.
        if result['customization']['formula_cells']:
            errors.append('formula_valued_delta_requires_review')
        if result['customization']['blank_cells'] and not result['customization']['formula_cells']:
            warnings.append('literal_blank_deltas_normalized_to_zero')
        result['customization']['has_nonzero_deltas'] = any(v not in (None, 0) for p in result['customization']['delta_paths'].values() for v in p)
        for name, address in TERMS.items():
            f, v = formulas[EXTERNAL][address], cached[EXTERNAL][address].value
            result['customization']['terms'][name] = {'mode': 'formula' if f.data_type == 'f' else 'value',
                                                     'value': float(v) if _number(v) else None}
            if not _number(v):
                errors.append(f'term_cached_value_missing:{name}')
            label = {'rate': 'average interest rate', 'grace_years': 'average grace period',
                     'maturity_years': 'average maturity'}[name]
            if not _contains(cached[EXTERNAL]['B' + address[1:]].value, label):
                errors.append(f'term_anchor_invalid:{name}')
        for metric, title, units, sheet, baseline, custom, col, header in METRICS:
            labelcol = 1 if col == 3 else 2
            if not _contains(cached[sheet].cell(baseline, labelcol).value, 'baseline') or not _contains(cached[sheet].cell(custom, labelcol).value, 'a2.'):
                errors.append(f'output_anchor_invalid:{metric}')
            actual = [cached[sheet].cell(header, col + i).value for i in range(21)]
            if actual != years:
                errors.append(f'output_year_geometry_invalid:{metric}')
            if any(formulas[sheet].cell(r, col + i).data_type != 'f'
                   for r in (baseline, custom) for i in range(21)):
                errors.append(f'output_formula_missing:{metric}')
        missing, error_values = 0, 0
        for probe in probes():
            value = cached[probe['sheet']][probe['cell']].value
            missing += int(value is None)
            error_values += int(isinstance(value, str) and value.startswith('#'))
        result['evidence']['saved_cache'] = 'present' if not missing else 'incomplete'
        result['evidence']['cached_probe_count'] = len(probes())
        result['evidence']['missing_cached_probes'] = missing
        result['evidence']['cached_error_probes'] = error_values
        if missing:
            errors.append('output_cache_incomplete')
        if error_values:
            warnings.append('cached_output_errors_present')
        if 'CI Summary' in cached.sheetnames:
            ws = cached['CI Summary']
            for key, address, anchor, label in THRESHOLDS:
                value = ws[address].value
                if _contains(ws['B1'].value, 'debt carrying capacity and thresholds') and _contains(ws[anchor].value, label) and _number(value):
                    result['thresholds'][key] = float(value)
                else:
                    result['thresholds'][key] = None
                    warnings.append(f'threshold_unavailable:{key}')
        else:
            warnings.append('thresholds_unavailable')
        result['compatibility']['checks'] = ['input_row_anchors', 'input_years', 'additive_input_formulas', 'term_anchors', 'output_anchors', 'output_years', 'output_formula_presence', 'cached_values']
        result['compatibility']['status'] = 'supported' if not errors else 'unsupported'
        result['evidence']['structural'] = 'passed' if not errors else 'failed'
        warnings.append('formula_changes_not_certified_by_structure')
        return result
    finally:
        formulas.close()
        cached.close()
        if file_hash(path) != original_hash:
            raise IntakeError('workbook_changed_during_inspection')
