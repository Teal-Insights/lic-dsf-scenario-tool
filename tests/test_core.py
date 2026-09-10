"""Neutral contract and spreadsheet-semantics regressions; no workbook fixtures."""
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from lic_dsf.compare import compare_one
from lic_dsf.engine import digest, _Graph
from lic_dsf.ida21 import DELTA_ROWS, PUBLIC, EXTERNAL, cell
from lic_dsf.scenarios import InputError, zero_scenario, normalize_scenario, imported_scenario
from lic_dsf.workbooks import IntakeError, check_archive, _additive_input


class Contracts(unittest.TestCase):
    def setUp(self):
        self.info = {'input_years': list(range(2030, 2051))}
        self.case = zero_scenario(self.info)
    def test_full_horizon_and_hash(self):
        self.case['delta_paths']['20'][20] = -0.125
        normalized = normalize_scenario(self.case, self.info['input_years'])
        expected = hashlib.sha256(json.dumps(normalized, sort_keys=True, separators=(',', ':'), ensure_ascii=False, allow_nan=False).encode()).hexdigest()
        self.assertEqual(digest(normalized), expected)
        self.assertEqual(sum(map(len, normalized['delta_paths'].values())), 273)
        self.assertIsNone(normalized['terms'])
    def test_incomplete_path_rejected(self):
        self.case['delta_paths']['20'].pop()
        with self.assertRaises(InputError): normalize_scenario(self.case, self.info['input_years'])
    def test_missing_row_rejected(self):
        del self.case['delta_paths']['38']
        with self.assertRaises(InputError): normalize_scenario(self.case, self.info['input_years'])
    def test_ui_blank_bool_nonfinite_rejected(self):
        for value in (None, '', True, float('nan'), float('inf')):
            self.case['delta_paths']['20'][0] = value
            with self.assertRaises(InputError): normalize_scenario(self.case, self.info['input_years'])
    def test_term_inheritance_and_constraints(self):
        for terms in ({'rate': .05}, {'rate':.05,'grace_years':5,'maturity_years':5}, {'rate':.05,'grace_years':1.5,'maturity_years':9}):
            self.case['terms'] = terms
            with self.assertRaises(InputError): normalize_scenario(self.case, self.info['input_years'])
        self.case['terms'] = {'rate': .05, 'grace_years': 2, 'maturity_years': 9}
        self.assertEqual(normalize_scenario(self.case,self.info['input_years'])['terms']['grace_years'],2.0)
    def test_import_preserves_every_nonzero(self):
        for i, row in enumerate(DELTA_ROWS): self.case['delta_paths'][str(row)][20-i] = i+.5
        self.info['customization'] = {'delta_paths':self.case['delta_paths']}
        self.assertEqual(imported_scenario(self.info), self.case)
    def test_literal_blank_requires_direct_addition(self):
        self.assertTrue(_additive_input("='Reference'!B2+$E$20",'E20'))
        self.assertTrue(_additive_input("='Reference'!B2-+'Reference'!B3+E20", 'E20'))
        for formula in ('=SUM(E20)', '=IF(E20="",1,0)', '=IF(E20="",1,0)+E20',
                        '=ISBLANK(E20)+E20', '=SUM(B2)+E20', '=B2+E21', None):
            self.assertFalse(_additive_input(formula,'E20'))
    def test_strict_comparator(self):
        for expected, actual in [(0,None),(0,False),(1,'1'),('#N/A','#VALUE!'),(float('nan'),float('nan'))]:
            self.assertFalse(compare_one(expected,actual)[0])
        self.assertTrue(compare_one(1,1.0000001)[0])
        self.assertTrue(compare_one('#VALUE!','#VALUE!')[0])
    def test_unsafe_archive_and_entities_rejected(self):
        for name, contents in [('../escape.xml','x'), ('xl/a.xml','<!DOCTYPE z>'), ('xl/a.xml','<!ENTITY z>')]:
            with tempfile.TemporaryDirectory() as td:
                path=Path(td)/'neutral.xlsx'
                with zipfile.ZipFile(path,'w') as z: z.writestr(name,contents)
                with self.assertRaises(IntakeError): check_archive(path)


class Semantics(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from lic_dsf.excel_semantics import apply
        apply()
    def test_empty_cell_is_zero_empty_text_is_error(self):
        from excel_grapher.core import operators
        from excel_grapher.core.types import XlError
        self.assertEqual(operators.to_number(None),0)
        self.assertEqual(operators.to_number(''),XlError.VALUE)
    def test_range_text_ignored_literal_text_coerced(self):
        from excel_grapher.evaluator.functions import FUNCTIONS
        self.assertEqual(FUNCTIONS['SUM']([2,'7',True,None]),2)
        self.assertEqual(FUNCTIONS['SUM'](2,'7',True),10)
    def test_aggregate_errors_and_sample_stdev(self):
        from excel_grapher.evaluator.functions import FUNCTIONS
        from excel_grapher.core.types import XlError
        self.assertEqual(FUNCTIONS['SUM']([2,XlError.NA]),XlError.NA)
        self.assertAlmostEqual(FUNCTIONS['STDEV']([1,3]),2**.5)
        self.assertEqual(FUNCTIONS['STDEV']([1]),XlError.DIV)
    def test_term_formula_restore_and_full_surface(self):
        from excel_grapher import create_dependency_graph
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'neutral.xlsx'; w=openpyxl.Workbook();w.active.title=PUBLIC;s=w.create_sheet(EXTERNAL)
            for row in DELTA_ROWS:
                for col in range(5,26):w[PUBLIC].cell(row,col,0)
            s['C21']='=0.05';s['C23']='=2';s['C24']='=9';w.save(p);w.close()
            g=_Graph.__new__(_Graph);g.graph=create_dependency_graph(p,[f"'{PUBLIC}'!E10:Y38",f"'{EXTERNAL}'!C21:C24"],load_values=True);g.keys={str(k).upper().replace("'",''):k for k in g.graph.keys()};g.original={}
            key=g.keys[f'{EXTERNAL}!C21'.upper()];original=g.graph.get_node(key).formula
            case=zero_scenario({'input_years':list(range(21))});case['delta_paths']['38'][20]=.25
            g.apply(case);self.assertEqual(g.graph.get_node(key).formula,original)
            case['terms']={'rate':.07,'grace_years':3,'maturity_years':10};g.apply(case)
            self.assertIsNone(g.graph.get_node(key).formula);g.restore();self.assertEqual(g.graph.get_node(key).formula,original)
            self.assertEqual(g.graph.get_node(g.keys[f'{PUBLIC}!Y38'.upper()]).value,0)
            del g.keys[f'{PUBLIC}!Y38'.upper()]
            with self.assertRaises(IntakeError): g.apply(case)
            self.assertEqual(g.original,{})

if __name__ == '__main__': unittest.main()
