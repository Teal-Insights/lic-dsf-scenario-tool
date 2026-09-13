import copy, csv, io, json, unittest
from unittest.mock import patch
from zipfile import ZipFile
from hashlib import sha256
from openpyxl import load_workbook
from test_charts import fixture
from lic_dsf.app import shareable
from lic_dsf.briefing_pack import tables, workbook_bytes, render_pack, _csv
from lic_dsf.charts import render_indicator


def public_fixture():
    c=fixture();c['chart_context']={'workbook_sha256':'a'*64,'label':None,'revision':0}
    for run in c['runs']:
        run['revision']=1;run['private_journal']='DO_NOT_SHARE_829'
        r=run['result'];r.update(contract_version='test',scenario_hash='c'*64,first_projection_year=2024,input_years=[2024,2025,2026])
        r['engine_identity']={'core_sha256':'d'*64};run['shared_rationale']={'20':'=A1+1'}
    return shareable(c)


class BriefingPackTest(unittest.TestCase):
    def test_tables_preserve_values_years_units_and_missing_states(self):
        c=public_fixture();c['runs'][0]['result']['points'][0]['scenario']=None
        c['runs'][0]['result']['points'][1]['scenario']='#N/A'
        before=copy.deepcopy(c);t=tables(c)
        self.assertEqual(len(t['Results']),109)
        self.assertIsNone(t['Results'][1][8]);self.assertIsNone(t['Results'][1][9])
        self.assertEqual(t['Results'][4][8],'#N/A')
        self.assertEqual([r[2] for r in t['Macro assumptions'][1:4]],[2024,2025,2026])
        self.assertEqual(t['Macro assumptions'][2][3],-.1)
        self.assertEqual(c,before)
        with self.assertRaisesRegex(ValueError,'align'):
            c['runs'][0]['result']['input_years']=[2024];tables(c)

    def test_excel_and_csv_never_execute_user_text(self):
        t=tables(public_fixture());wb=load_workbook(io.BytesIO(workbook_bytes(t)))
        self.assertEqual(wb.sheetnames[0],'Read me first')
        self.assertIn('What this analysis answers',[row[0].value for row in wb['Read me first']])
        self.assertFalse(wb._external_links)
        self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row))
        self.assertEqual(wb['Shared explanations']['D2'].value,'=A1+1')
        self.assertEqual(wb['Shared explanations']['D2'].data_type,'s')
        for text in ['=1+1','+SUM(A1)','-1+2','@SUM(1)',' \t=1']:
            result=list(csv.reader(io.StringIO(_csv([[text,-.1,None]]).decode('utf-8-sig'))))[0]
            self.assertTrue(result[0].startswith("'"));self.assertEqual(result[1],'-0.1');self.assertEqual(result[2],'')

    def test_packet_inventory_hashes_and_private_exclusions(self):
        c=public_fixture();before=copy.deepcopy(c)
        with patch('lic_dsf.briefing_pack.render_comparison',return_value=b'PDF'),patch('lic_dsf.briefing_pack.render_indicator',return_value=b'CHART'):
            blob=render_pack(c)
        with ZipFile(io.BytesIO(blob)) as z:
            self.assertEqual(len(z.namelist()),35)
            manifest=json.loads(z.read('manifest.json'))
            for name,entry in manifest['files'].items():
                self.assertFalse(name.startswith('/'));self.assertNotIn('..',name.split('/'))
                self.assertEqual(sha256(z.read(name)).hexdigest(),entry['sha256'])
            self.assertEqual(json.loads(z.read('data-and-assumptions.json')),c)
            for name in z.namelist():
                self.assertNotIn(b'DO_NOT_SHARE_829',z.read(name));self.assertNotIn(b'PRIVATE_SENTINEL_X92',z.read(name))
            self.assertEqual(len([n for n in z.namelist() if n.endswith('.png')]),12)
            self.assertEqual(len([n for n in z.namelist() if n.endswith('.svg')]),12)
        self.assertEqual(c,before)

    def test_individual_svg_portable_and_rejects_unknown_metric(self):
        c=public_fixture();before=copy.deepcopy(c)
        for view in ['standard','briefing']:
            svg=render_indicator(c,'public_pv_gdp',view,'svg')
            self.assertIn(b'<svg',svg);self.assertNotIn(b'<script',svg);self.assertNotIn(b'<foreignObject',svg)
            self.assertNotIn(b'PRIVATE_SENTINEL_X92',svg)
        for metric in ['../private',None,{},'unknown']:
            with self.assertRaises(ValueError):render_indicator(c,metric)
        self.assertEqual(c,before)

    def test_exact_bundled_fonts_and_unsupported_scenario_labels(self):
        from pathlib import Path
        from matplotlib.font_manager import FontProperties, findfont
        from lic_dsf.typography import FONT_FOLDER, FONT_FAMILY, font_properties
        from lic_dsf.charts import render_comparison
        for weight, name in [('normal','Inter-Regular.otf'), ('bold','Inter-SemiBold.otf')]:
            self.assertEqual(Path(findfont(FontProperties(family=FONT_FAMILY, weight=weight))), FONT_FOLDER/name)
            self.assertEqual(Path(font_properties(weight=weight).get_file()), FONT_FOLDER/name)
        self.assertEqual(Path(font_properties(family='IBM Plex Serif', weight='bold').get_file()), FONT_FOLDER/'IBMPlexSerif-SemiBold.otf')
        c=public_fixture();c['runs'][0]['share_label']='Scenario Կ'
        for view in ['standard','briefing']:
            with self.assertRaisesRegex(ValueError,'export fonts bundled'):
                render_indicator(c,'public_pv_gdp',view,'svg')
            with self.assertRaisesRegex(ValueError,'export fonts bundled'):
                render_comparison(c,view,'pdf')
        # Data-only exports retain unsupported characters without executing them.
        self.assertIn('Scenario Կ',str(tables(c)['Results']))

    def test_financing_guide_does_not_invent_an_alternative(self):
        from lic_dsf.reader_guide import financing_comparison_note
        self.assertIn('same financing',financing_comparison_note([{'rate':.08},{'rate':.08}]))
        self.assertIn('differs',financing_comparison_note([{'rate':.04},{'rate':.08}]))
        self.assertIn('Only one',financing_comparison_note([{'rate':.08}]))

    def test_table_only_export_names_source_comparison_and_guidance(self):
        c=public_fixture();wb=load_workbook(io.BytesIO(workbook_bytes(tables(c), c)))
        guide={row[0].value:row[1].value for row in wb['Read me first'] if row[0].value}
        self.assertEqual(guide['Comparison case'],'Scenario B')
        self.assertIn('Analyst-supplied',guide['Starting source'])
        self.assertIn('pending',guide['Recorded check status'])
        self.assertIn('p. 37',guide['Illustrative financing scenarios'])
        links=[cell.hyperlink.target for row in wb['Read me first'] for cell in row if cell.hyperlink]
        self.assertTrue(any('#page=38' in url for url in links))
